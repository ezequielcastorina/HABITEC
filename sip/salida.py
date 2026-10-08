"""Salidas no gráficas: planilla de paneles (xlsx) e informe de validación (txt)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C
from .contorno import texto_caida
from .modelo import Despiece


def planilla_xlsx(d: Despiece, ruta: Path, car=None, fecha: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Paneles"
    ws["A1"] = "HABITEC · Proyecto"
    ws["B1"] = d.plano.proyecto
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    ws["A2"] = f"Caída hacia el lado {texto_caida(d.plano)}"
    ws["C1"] = f"Emisión {fecha}" if fecha else ""
    if car is not None:
        ws["F1"] = "Rev. interior"
        ws["G1"] = car.interior or "a definir"
        ws["F2"] = "Cielorraso"
        ws["G2"] = car.cielorraso or "a definir"
        ws["I1"] = "Rev. exterior"
        ws["J1"] = car.exterior or "a definir"
        ws["I2"] = "Revisión"
        ws["J2"] = car.revision
    cab = ["Código", "Tipo", "Lado", "N°", "Ancho (mm)", "Alto/Largo (mm)", "Espesor (mm)", "Ajuste",
           "Revestimiento", "Cara", "Sentido", "Vanos", "Rebaje perimetral", "Tirantes y tapas", "Alertas"]
    fila0 = 4
    for j, h in enumerate(cab, start=1):
        c = ws.cell(row=fila0, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = fila0 + 1
    for p in d.muros:
        ws.append([])
        vanos = "; ".join(f"{v.vano_id} ({C.TIPOS_VANO[v.tipo]}"
                          f"{'' if v.completo else ', compartido'})" for v in p.vanos)
        vals = [p.codigo, "Muro", p.lado, p.numero, round(p.largo * 1000), round(p.alto * 1000),
                round(p.espesor * 1000), "sí" if p.es_ajuste else "",
                "Smart panel" if p.smart else "No", p.cara_rev, p.sentido_rev, vanos,
                "30 mm, 4 lados", _bordes_txt(p), "; ".join(p.alertas)]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    for t in d.techos:
        vals = [t.codigo, "Techo", "", t.numero, round(t.ancho * 1000), round(t.largo * 1000),
                round(C.ESPESOR_PANEL * 1000), "sí" if t.es_ajuste else "", "", "", "", "",
                ((f"30 mm, 3 lados (sin lado {t.caida_hacia})" if t.en_caida else "30 mm, 4 lados")
                 ), _tirantes_techo(t), "; ".join(t.alertas)]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    # Filas con alertas resaltadas
    rojo = PatternFill("solid", fgColor="F8D7D3")
    for fila in ws.iter_rows(min_row=fila0 + 1, max_row=r - 1):
        if fila[-1].value:
            for c in fila:
                c.fill = rojo
    # Solapa con la secuencia de carga del camión
    from .carga import nombre_esquina, secuencia, vertice_arranque
    items, _v0 = secuencia(d, getattr(car, "esquina", "NO"), getattr(car, "parrillas", "no") == "si")
    wc = wb.create_sheet("Carga camión")
    wc.cell(row=1, column=1, value=f"{d.plano.proyecto} · secuencia de carga del camión").font = Font(bold=True, size=12)
    wc.cell(row=2, column=1, value=f"Esquina de arranque del montaje: {nombre_esquina(d, _v0)}. "
                                   "El orden de carga es el orden de montaje; el techo va último.")
    for j, h in enumerate(("Orden", "Grupo", "Código", "Medidas (mm)", "Observaciones"), start=1):
        c = wc.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
    for n, it in enumerate(items, start=1):
        for j, v in enumerate((n, it.grupo, it.codigo, it.medidas, it.detalle), start=1):
            wc.cell(row=4 + n, column=j, value=v)
    for j, a in enumerate((8, 12, 18, 16, 20), start=1):
        wc.column_dimensions[get_column_letter(j)].width = a
    anchos = [9, 8, 6, 5, 11, 15, 12, 8, 14, 10, 11, 36, 28, 50, 60]
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=1)
    wb.save(ruta)


def informe_txt(d: Despiece, ruta: Path, car=None) -> None:
    L = [f"HABITEC · INFORME DE VALIDACIÓN — {d.plano.proyecto}", "=" * 60,
         f"Módulo: {d.plano.modulo.ancho_x:.3f} × {d.plano.modulo.alto_y:.3f} m",
         f"Caída hacia el lado: {texto_caida(d.plano) if d.plano.caida_hacia else 'no definida'}",
         f"Paneles de muro: {len(d.muros)} · Paneles de techo: {len(d.techos)} · Vanos: {len(d.plano.vanos)}"]
    if car is not None:
        L += [f"Revestimiento interior: {car.interior or 'a definir'}",
              f"Cielorraso: {car.cielorraso or 'a definir'}",
              f"Revestimiento exterior: {car.exterior or 'a definir'}"]
    L.append("")
    if d.avisos:
        L.append("AVISOS (no bloquean):")
        L += [f"  - {a}" for a in d.avisos]
        L.append("")
    if d.alertas:
        L.append(f"ALERTAS ({len(d.alertas)}):")
        L += [f"  - {a}" for a in d.alertas]
    else:
        L.append("Sin alertas: el dibujo cumple todas las reglas verificadas.")
    ruta.write_text("\n".join(L) + "\n", encoding="utf-8")


def _bordes_txt(p) -> str:
    nombres = {"izq": "izq.", "der": "der."}
    partes = []
    for b, (clase, dato) in p.bordes.items():
        if clase == "tirante":
            partes.append(f"{nombres[b]}: trae tirante 50x70 (junta con {dato})")
        elif clase == "recibe":
            partes.append(f"{nombres[b]}: rebaje libre (recibe de {dato})")
        else:
            partes.append(f"{nombres[b]}: tapa 25x70 (esquina {p.lado}-{dato})")
    for jj in p.jambas_junta:
        partes.append(f"{nombres[jj.borde]}: trae tirante 50x70 de la jamba de {jj.vano_id} (vano de {jj.con_panel})")
    return "; ".join(partes)


def _tirantes_techo(t) -> str:
    lleva = [f"{'+'.join(v)}" for l, v in t.uniones.items() if t.tirantes.get(l) == "lleva"]
    recibe = [f"{'+'.join(v)}" for l, v in t.uniones.items() if t.tirantes.get(l) == "recibe"]
    partes = []
    if lleva:
        partes.append("trae tirante 50x70 hacia " + ", ".join(lleva))
    if recibe:
        partes.append("recibe tirante de " + ", ".join(recibe))
    return "; ".join(partes)
